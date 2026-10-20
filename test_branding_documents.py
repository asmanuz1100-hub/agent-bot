"""Kompaniya rekvizitlari akt sverka hujjatlarida: kiritilmasa — avvalgi nom, kiritilsa — mijoz brendi."""
import io
import os
import tempfile
import unittest

import branding
import core
import reports


class BrandingDocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = core.connect(os.path.join(self.tmp.name, 'b.sqlite3'))
        self.db.executemany('INSERT INTO users(id,role,name) VALUES(?,?,?)', [(1, 'admin', 'R'), (2, 'agent', 'Ali')])
        self.db.execute("INSERT INTO clients(id,agent,name,shop_name,phone,address) VALUES(1,2,'Mijoz','Do‘kon','+998','Ko‘cha')")
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def _xlsx_title(self, raw):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw)).active
        return ws['A1'].value, ws['A2'].value

    def test_default_keeps_previous_name(self):
        r = reports.reconciliation(self.db, 1, 1)
        self.assertEqual(self._xlsx_title(reports.reconciliation_xlsx(r))[0], 'ASMAN SILICAT — МИЖОЗ АКТ СВЕРКА')
        self.assertEqual(branding.file_prefix(self.db), 'ASMAN')

    def test_client_brand_and_details_are_used(self):
        branding.save(self.db, 'Namuna Savdo', None, {'legalName': 'Namuna Savdo MChJ', 'address': 'Toshkent', 'inn': '123456789', 'director': 'A. Karimov'})
        r = reports.reconciliation(self.db, 1, 1)
        title, line = self._xlsx_title(reports.reconciliation_xlsx(r))
        self.assertEqual(title, 'Namuna Savdo MChJ — МИЖОЗ АКТ СВЕРКА')
        self.assertIn('STIR (INN): 123456789', line)
        self.assertTrue(reports.reconciliation_pdf(r).startswith(b'%PDF'))
        self.assertTrue(reports.all_clients_pdf(self.db, 1).startswith(b'%PDF'))
        self.assertEqual(self._xlsx_title(reports.all_clients_xlsx(self.db, 1))[0], 'Namuna Savdo MChJ — БАРЧА МИЖОЗЛАР АКТ СВЕРКА')
        self.assertEqual(branding.file_prefix(self.db), 'Namuna-Savdo')

    def test_details_validation(self):
        with self.assertRaises(ValueError):
            branding.save(self.db, details_in={'inn': '1' * 30})


if __name__ == '__main__':
    unittest.main()
